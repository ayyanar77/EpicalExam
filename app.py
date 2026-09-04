import os
import io
from flask import Flask, render_template, request, redirect, url_for, session, flash, send_file
from dotenv import load_dotenv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from sheets import (
    verify_candidate,
    get_questions,
    save_exam_result,
    get_all_results,
    get_gspread_client
)

# Load environment variables
load_dotenv()

app = Flask(__name__)
app.secret_key = os.getenv("SECRET_KEY", "epical_secret_key_2026")

# Admin credentials from environment variables
ADMIN_USERNAME = os.getenv("ADMIN_USERNAME", "admin")
ADMIN_PASSWORD = os.getenv("ADMIN_PASSWORD", "admin123")


@app.route("/")
def index():
    if "user_role" in session:
        if session["user_role"] == "admin":
            return redirect(url_for("admin_dashboard"))
        elif session["user_role"] == "candidate":
            return redirect(url_for("instructions"))
    return redirect(url_for("login"))


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        login_type = request.form.get("login_type")
        
        if login_type == "admin":
            username = request.form.get("username", "").strip()
            password = request.form.get("password", "")
            
            if username == ADMIN_USERNAME and password == ADMIN_PASSWORD:
                session["user_role"] = "admin"
                session["user_name"] = "Administrator"
                session["user_email"] = "admin@epicalexam.com"
                flash("Admin logged in successfully", "success")
                return redirect(url_for("admin_dashboard"))
            else:
                flash("Invalid admin credentials", "danger")
                
        elif login_type == "candidate":
            email = request.form.get("email", "").strip()
            password = request.form.get("password", "")
            
            candidate_data = verify_candidate(email, password)
            if candidate_data:
                session["user_role"] = "candidate"
                session["user_name"] = candidate_data.get("name", "Candidate")
                session["user_email"] = candidate_data.get("email", email)
                session["roll_number"] = candidate_data.get("roll_number", "N/A")
                flash(f"Welcome {session['user_name']}!", "success")
                return redirect(url_for("instructions"))
            else:
                flash("Invalid candidate email or password", "danger")

    return render_template("login.html")


@app.route("/logout")
def logout():
    session.clear()
    flash("You have been logged out successfully.", "info")
    return redirect(url_for("login"))


@app.route("/instructions")
def instructions():
    if session.get("user_role") != "candidate":
        flash("Please log in as a candidate to view instructions.", "danger")
        return redirect(url_for("login"))
    return render_template("instructions.html")


@app.route("/exam")
def exam():
    if session.get("user_role") != "candidate":
        flash("Please log in as a candidate to take the exam.", "danger")
        return redirect(url_for("login"))
        
    questions = get_questions()
    return render_template("exam.html", questions=questions)


@app.route("/submit_exam", methods=["POST"])
def submit_exam():
    if session.get("user_role") != "candidate":
        flash("Unauthorized submission.", "danger")
        return redirect(url_for("login"))
        
    questions = get_questions()
    total_questions = len(questions)
    score = 0
    
    # Server-side grading
    for q in questions:
        q_id = q["id"]
        submitted_answer = request.form.get(f"q_{q_id}")
        correct_answer = q.get("correct_option", "").strip().upper()
        
        if submitted_answer and submitted_answer.strip().upper() == correct_answer:
            score += 1
            
    percentage = (score / total_questions * 100) if total_questions > 0 else 0
    status = "PASSED" if percentage >= 50 else "FAILED"
    
    candidate_name = session.get("user_name", "Candidate")
    email = session.get("user_email", "")
    roll_number = session.get("roll_number", "N/A")
    
    # Save to database (Google Sheets and backup)
    save_exam_result(candidate_name, email, roll_number, score, total_questions, percentage, status)
    
    # Store latest result in session for display
    session["latest_result"] = {
        "Name": candidate_name,
        "Email": email,
        "RollNumber": roll_number,
        "Score": score,
        "Total": total_questions,
        "Percentage": f"{percentage:.2f}%",
        "Status": status
    }
    
    flash("Exam submitted successfully!", "success")
    return redirect(url_for("result"))


@app.route("/result")
def result():
    if session.get("user_role") != "candidate":
        return redirect(url_for("login"))
        
    latest_result = session.get("latest_result")
    if not latest_result:
        flash("No recent exam result found.", "info")
        return redirect(url_for("instructions"))
        
    # Get all results for leaderboard
    all_results = get_all_results()
    
    # Sort by score (descending)
    def parse_score(item):
        try:
            return float(item.get("Score", 0))
        except (ValueError, TypeError):
            return 0
            
    leaderboard = sorted(all_results, key=parse_score, reverse=True)
    
    try:
        pct_value = float(str(latest_result.get("Percentage", "0")).replace("%", ""))
    except ValueError:
        pct_value = 0
        
    return render_template(
        "result.html",
        result=latest_result,
        percentage=pct_value,
        leaderboard=leaderboard
    )


@app.route("/admin_dashboard")
def admin_dashboard():
    if session.get("user_role") != "admin":
        flash("Unauthorized access. Please log in as admin.", "danger")
        return redirect(url_for("login"))
        
    results = get_all_results()
    passed_count = sum(1 for r in results if str(r.get("Status", "")).upper() == "PASSED")
    failed_count = sum(1 for r in results if str(r.get("Status", "")).upper() == "FAILED")
    
    _, sheet = get_gspread_client()
    is_sheets_connected = (sheet is not None)
    
    return render_template(
        "admin_dashboard.html",
        results=results,
        passed_count=passed_count,
        failed_count=failed_count,
        is_sheets_connected=is_sheets_connected
    )


@app.route("/export_excel")
def export_excel():
    if session.get("user_role") != "admin":
        flash("Unauthorized access.", "danger")
        return redirect(url_for("login"))
        
    results = get_all_results()
    
    # Create workbook using openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Exam Results"
    
    headers = ["Timestamp", "Roll Number", "Candidate Name", "Email", "Score", "Total", "Percentage", "Status"]
    ws.append(headers)
    
    # Style header row
    header_fill = PatternFill(start_color="0D6EFD", end_color="0D6EFD", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        
    for r in results:
        ws.append([
            r.get("Timestamp", "-"),
            r.get("RollNumber", "-"),
            r.get("Name", "-"),
            r.get("Email", "-"),
            r.get("Score", 0),
            r.get("Total", 0),
            r.get("Percentage", "0%"),
            r.get("Status", "-")
        ])
        
    # Auto adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
        
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="EpicalExam_Results.xlsx"
    )


if __name__ == "__main__":
    app.run(debug=True, port=5000)
